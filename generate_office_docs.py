"""
Script to generate editable Microsoft Word (.docx) project report
and Microsoft PowerPoint (.pptx) evaluation presentation.
Uses python-docx and python-pptx libraries installed in the venv.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from pptx.dml.color import RGBColor as PptxRGBColor


def generate_word_report(output_path):
    """Generates the complete 28-section project report in Microsoft Word format."""
    doc = Document()

    # Define color scheme
    primary_color = RGBColor(37, 99, 235)  # Royal Blue
    dark_color = RGBColor(15, 23, 42)      # Slate Dark
    gray_color = RGBColor(100, 116, 139)   # Muted Gray

    # --- Title Page ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("IT HELP DESK TICKET MANAGEMENT SYSTEM\n")
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = primary_color

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("A Project Report Submitted in Partial Fulfillment of the Requirements\nfor the Degree of Bachelor of Engineering / Technology\n")
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = gray_color

    doc.add_paragraph("\n" * 3)

    student_p = doc.add_paragraph()
    student_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st_run = student_p.add_run("Submitted by:\n[STUDENT FULL NAME]\nRoll No. / Enrollment ID: [STUDENT ROLL NUMBER]\nAcademic Year: 2025 – 2026\n")
    st_run.font.size = Pt(12)
    st_run.font.bold = True

    doc.add_paragraph("\n" * 2)

    guide_p = doc.add_paragraph()
    guide_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    gd_run = guide_p.add_run("Under the Guidance of:\nInternal Academic Mentor: [FACULTY GUIDE NAME]\nExternal Industry Supervisor: [INDUSTRY SUPERVISOR NAME]\n\nDepartment of Computer Engineering / Information Technology\n[COLLEGE / UNIVERSITY NAME]\n[CAMPUS LOCATION]\n")
    gd_run.font.size = Pt(11)

    doc.add_page_break()

    # --- Content Sections ---
    sections = [
        ("1. Cover Page", "See title page above."),
        ("2. Student Details and Internship Details", 
         "Student Name: [STUDENT FULL NAME]\n"
         "Roll Number / PRN: [STUDENT ROLL NUMBER]\n"
         "Branch: Computer Engineering / Information Technology\n"
         "Semester: 6th Semester / 3rd Year\n"
         "College: [COLLEGE / UNIVERSITY NAME]\n"
         "Internship Organization: [NAME OF INTERNSHIP COMPANY]\n"
         "Internship Role: Web and Software Applications Intern\n"
         "Internship Duration: [START DATE] to [END DATE]\n"
         "Supervisor: [SUPERVISOR NAME, DESIGNATION]"),
        
        ("3. Certificate of Internship Evaluation", 
         "This is to certify that the project report entitled 'IT Help Desk Ticket Management System' is a bonafide work carried out by [STUDENT FULL NAME] (Roll No: [STUDENT ROLL NUMBER]) in partial fulfillment of the academic requirements for the internship evaluation in the Bachelor of Engineering curriculum during the academic year 2025–2026.\n\n"
         "[FACULTY GUIDE NAME] (Internal Project Mentor)    |    [HEAD OF DEPARTMENT]\n"
         "External Examiner                                 |    Date of Evaluation: ____________"),
        
        ("4. Declaration", 
         "I hereby declare that this project report titled 'IT Help Desk Ticket Management System' is an authentic record of my own work performed during my internship in Web and Software Applications. Any libraries or methodologies utilized have been duly referenced.\n\n"
         "[STUDENT FULL NAME]\nRoll Number: [STUDENT ROLL NUMBER]"),
        
        ("5. Acknowledgement", 
         "I express my deepest gratitude to the faculty and management of [COLLEGE NAME] and my internal project guide, [FACULTY GUIDE NAME]. I also thank my industry supervisor, [SUPERVISOR NAME], and colleagues at [INTERNSHIP ORGANIZATION] for their guidance in web development and software testing best practices."),
        
        ("6. Abstract", 
         "In contemporary educational institutions, reliable IT infrastructure is vital. However, technical complaints regarding workstations, network connectivity, software packages, and peripherals are frequently reported through unstructured communication channels, leading to forgotten requests and lack of accountability.\n\n"
         "This project presents the IT Help Desk Ticket Management System, a full-stack web application developed using Python with Flask, an SQLite relational database, and an HTML5/CSS3/JavaScript frontend. It enables users to report complaints across 6 IT categories, auto-generates sequential ticket IDs (e.g. HD-1001), provides an operational dashboard displaying dynamic statistics, and equips support staff with status management and activity logging controls. The application implements parameter-driven queries, CSRF mitigation, and automated test coverage, running locally on macOS without paid APIs."),
        
        ("7. Table of Contents", 
         "1. Cover Page\n2. Student Details\n3. Certificate\n4. Declaration\n5. Acknowledgement\n6. Abstract\n7. Table of Contents\n8. Introduction\n9. Problem Statement\n10. Objectives\n11. Existing System and Limitations\n12. Proposed System\n13. Software & Hardware Requirements\n14. Technology Stack\n15. System Architecture\n16. Functional and Non-Functional Requirements\n17. Use-Case Descriptions\n18. Data Flow Diagram\n19. Database Schema and ER Diagram\n20. Module Descriptions\n21. Implementation Details\n22. Testing and Test Cases\n23. Screenshots and Placeholders\n24. Results and Discussion\n25. Limitations\n26. Future Scope\n27. Conclusion\n28. References"),
        
        ("8. Introduction", 
         "Colleges and organizations maintain extensive IT hardware and software across multiple buildings. This project digitizes and automates the complaint management workflow, providing a transparent, audited channel between users and technicians."),
        
        ("9. Problem Statement", 
         "Manual and email-based IT support methods suffer from: loss of issue traceability, lack of priority triage, absence of accountability, and zero operational analytics."),
        
        ("10. Objectives", 
         "- Centralize technical complaint registration.\n- Automatically assign unique, human-readable ticket IDs (HD-1001).\n- Provide real-time dashboard analytics calculated from SQLite.\n- Enable multi-parameter search and filtering.\n- Implement role-based status updating and activity logs.\n- Ensure zero-cost local execution on macOS with VS Code."),
        
        ("11. Existing System and Limitations", 
         "Existing approaches rely on verbal reports, paper logbooks, or unorganized email inboxes. These lack searchability, accountability, status tracking, and priority triage, leading to neglected issues and prolonged computer downtime."),
        
        ("12. Proposed System", 
         "The proposed application provides an open-source, full-stack solution with structured categorization, dynamic metrics, role-based staff authentication, and an audited activity timeline."),
        
        ("13. Software and Hardware Requirements", 
         "Software: macOS, VS Code, Python 3.10+, Flask 3.1+, SQLite 3, Git, HTML5, CSS3, JavaScript.\nHardware: Apple Silicon or Intel Core i5/i7, 4 GB RAM minimum, 500 MB free disk space."),
        
        ("14. Technology Stack", 
         "- Frontend: HTML5 semantic structure, CSS3 custom responsive design (CSS Grid/Flexbox), Vanilla JavaScript.\n- Backend: Python 3 with Flask WSGI micro-framework.\n- Database: SQLite 3 relational database engine.\n- Security: Werkzeug PBKDF2 password hashing, session CSRF tokens."),
        
        ("15. System Architecture", 
         "The system follows the Model-View-Controller (MVC) pattern. The browser client interacts via HTTP with Flask route controllers (app.py), which query the SQLite database (helpdesk.db) using parameterized SQL queries and render Jinja2 HTML templates."),
        
        ("16. Functional and Non-Functional Requirements", 
         "Functional: Ticket submission, unique ID generation, live metrics dashboard, search, multi-criteria filtering, staff authentication, lifecycle updates, activity logs.\nNon-Functional: Sub-second response time, parameterized queries against SQL injection, CSRF mitigation, responsive layout, modular architecture."),
        
        ("17. Use-Case Descriptions", 
         "Use Case 1 (User): Navigate to /tickets/new, fill details, submit ticket, receive generated ID HD-XXXX.\nUse Case 2 (Staff): Log in at /login, review ticket queue, inspect ticket, change status to In Progress / Resolved, append technician note."),
        
        ("18. Data Flow Diagram (DFD)", 
         "Level 0 Context: User/Staff submits requests to the Flask Web Application, which executes queries against the SQLite Database.\nLevel 1: Separates ticket submission, queue querying, credential validation, and status log updating into distinct processes."),
        
        ("19. Database Schema and ER Diagram", 
         "Entities: users, tickets, and ticket_notes.\nRelationship: tickets has a 1-to-many relationship with ticket_notes (Foreign Key on ticket_id with ON DELETE CASCADE). B-Tree indexes on ticket_id, status, priority, and created_at."),
        
        ("20. Module Descriptions", 
         "1. Dashboard: Aggregates metrics and category counts.\n2. Ticket Submission: Validates form inputs and generates IDs.\n3. Ticket Queue: Implements search and filtering.\n4. Ticket Details & Management: Shows details and updates status.\n5. Authentication: Handles login, logout, and hashed passwords."),
        
        ("21. Implementation Details", 
         "Highlights include the atomic unique ID generation algorithm (HD-XXXX), server-side CSRF token verification interceptor, and parameterized query execution."),
        
        ("22. Testing and Test Cases", 
         "18 automated test cases were executed via Python unittest against an isolated SQLite test database. All 18 tests passed (100% pass rate) with zero errors. Tested: startup, DB seeding, validation constraints, search, filtering, authentication, and CSRF protection."),
        
        ("23. Screenshots and Placeholders", 
         "[Insert Screenshot 1: IT Operations Dashboard with Live Metrics]\n"
         "[Insert Screenshot 2: Submit New IT Ticket Form with Urgency Cards]\n"
         "[Insert Screenshot 3: Searchable and Filterable Ticket Queue Table]\n"
         "[Insert Screenshot 4: Ticket Details View with Activity Log Timeline]\n"
         "[Insert Screenshot 5: Staff Login Portal with Credentials Notice]"),
        
        ("24. Results and Discussion", 
         "The application met all functional goals. All dashboard statistics update dynamically from live database records. Client and server validation prevent corrupt records. The interface renders cleanly on desktop and mobile."),
        
        ("25. Limitations", 
         "SQLite file locking limits concurrent write throughput under high traffic. Notifications are in-app only. File attachments (screenshots/dumps) are not yet supported."),
        
        ("26. Future Scope", 
         "Version 2.0 enhancements: SMTP automated email alerts, file upload support for logs/screenshots, SLA countdown timers, and NLP-assisted ticket triage."),
        
        ("27. Conclusion", 
         "The IT Help Desk Ticket Management System provides a complete, beginner-friendly, secure, and production-inspired solution for managing technical support requests within academic institutions. Built using Python, Flask, and SQLite, it fulfills all internship evaluation requirements."),
        
        ("28. References", 
         "1. Miguel Grinberg, 'Flask Web Development: Developing Web Applications with Python', O'Reilly Media.\n"
         "2. SQLite Development Team, 'SQLite Documentation', https://www.sqlite.org/docs.html\n"
         "3. Mozilla Developer Network (MDN), 'HTML5, CSS3, and JavaScript Guides', https://developer.mozilla.org/\n"
         "4. Werkzeug Utilities Documentation, 'Password Hashing & Secure Cookies', https://werkzeug.palletsprojects.com/\n"
         "5. Python Software Foundation, 'unittest & sqlite3 Standard Library', https://docs.python.org/3/")
    ]

    for title, content in sections:
        h = doc.add_heading(title, level=2)
        h.runs[0].font.color.rgb = primary_color
        p = doc.add_paragraph(content)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(10)

    doc.save(output_path)
    print(f"Successfully generated Word report: {output_path}")


def generate_pptx_presentation(output_path):
    """Generates the 10-slide PowerPoint presentation."""
    prs = Presentation()
    prs.slide_width = PptxInches(13.333)  # 16:9 Widescreen
    prs.slide_height = PptxInches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette
    c_primary = PptxRGBColor(37, 99, 235)   # Royal Blue
    c_dark = PptxRGBColor(15, 23, 42)       # Slate 900
    c_muted = PptxRGBColor(100, 116, 139)   # Slate 500
    c_bg = PptxRGBColor(248, 250, 252)      # Light Gray/Slate 50

    slides_data = [
        # Slide 1: Title Slide
        {
            "title": "IT Help Desk Ticket Management System",
            "subtitle": "Web & Software Applications Internship Project Evaluation\nAcademic Year 2025 – 2026",
            "bullets": [
                "Candidate Name: [STUDENT FULL NAME]   |   Roll Number: [STUDENT ROLL NUMBER]",
                "Branch: Computer Engineering / Information Technology",
                "College: [COLLEGE / UNIVERSITY NAME]",
                "Internship Organization: [NAME OF INTERNSHIP COMPANY]",
                "Guide / Mentor: [FACULTY GUIDE NAME]"
            ]
        },
        # Slide 2: Problem Statement
        {
            "title": "Introduction & Problem Statement",
            "subtitle": "Challenges in Institutional IT Complaint Handling",
            "bullets": [
                "Technical issues (lab crashes, network timeouts, printer jams) happen daily across campuses.",
                "Traditional reporting methods (verbal complaints, loose paper logs, unmonitored emails) cause lost requests.",
                "High-priority failures (e.g. server down) get delayed under routine requests.",
                "Zero accountability: No timestamped ownership or technician progress logs.",
                "Need: A centralized, transparent, and easy-to-use incident tracking web portal."
            ]
        },
        # Slide 3: Objectives
        {
            "title": "Project Objectives",
            "subtitle": "Core Goals Established for the Application",
            "bullets": [
                "Centralize all IT complaint submissions through an intuitive web form.",
                "Auto-generate unique, human-readable ticket tracking IDs (e.g., HD-1001).",
                "Real-time operational dashboard with dynamic statistics from SQLite.",
                "Multi-parameter search (by ID, title, requester) and filters (status, category, priority).",
                "Support staff authentication with role-based access control and activity logging.",
                "100% open-source local execution on macOS with zero paid services or API fees."
            ]
        },
        # Slide 4: Proposed System & Features
        {
            "title": "Proposed System & Key Features",
            "subtitle": "Core Functional Modules Implemented",
            "bullets": [
                "Operational Dashboard: Live count cards (Total, Open, In Progress, Resolved, High Priority).",
                "Category Breakdown: Visual distribution bars across 6 IT domains.",
                "Self-Service Submission: Input validation with error/success alerts.",
                "Searchable Ticket Queue: Filterable table displaying creation and update timestamps.",
                "Ticket Details & Lifecycle: Dedicated page with status transitions and technician notes timeline."
            ]
        },
        # Slide 5: Technology Stack
        {
            "title": "Technology Stack & Architecture",
            "subtitle": "Tools, Libraries, and Rationale",
            "bullets": [
                "Frontend: HTML5, modern CSS3 (Variables, Flexbox, Grid), Vanilla JavaScript (No heavy UI bloat).",
                "Backend: Python 3 with Flask (Lightweight, declarative WSGI routing, Jinja2 templating).",
                "Database: SQLite 3 (Self-contained, serverless relational database engine stored locally).",
                "Security: Werkzeug PBKDF2 password hashing, session-based CSRF protection.",
                "Environment & VCS: VS Code on macOS, Python Virtual Environment (venv), Git, GitHub."
            ]
        },
        # Slide 6: Architecture & Workflow
        {
            "title": "System Architecture & Workflow",
            "subtitle": "End-to-End Incident Resolution Lifecycle",
            "bullets": [
                "Architecture: Model-View-Controller (MVC) with clear separation of concerns.",
                "1. User Reports Issue: Validated form -> Unique ticket ID assigned -> Committed to SQLite.",
                "2. Live Triage: Dashboard metrics recalculate immediately via SQL aggregation queries.",
                "3. Staff Authentication: Technicians log in using securely hashed credentials.",
                "4. Status & Notes Update: Staff transitions status to In Progress / Resolved with audit logs."
            ]
        },
        # Slide 7: Database Design
        {
            "title": "Database Schema & Relationships",
            "bullets": [
                "users Table: Stores credentials (id, username, email, password_hash, role, created_at).",
                "tickets Table: Primary incident records (ticket_id, title, requester, category, priority, status).",
                "ticket_notes Table: Audited activity logs linked via Foreign Key (ticket_id ON DELETE CASCADE).",
                "Optimization: B-Tree database indexes on ticket_id, status, priority, and created_at.",
                "Security: Parameterized queries (cursor.execute(..., (?))) prevent SQL Injection."
            ]
        },
        # Slide 8: Implementation & Screens
        {
            "title": "Implementation & Key User Interface",
            "subtitle": "Key Responsive Screens Developed",
            "bullets": [
                "Dashboard View: High-level operations overview with dynamic summary cards.",
                "Create Ticket Form: User-friendly submission with priority selector radio cards.",
                "Ticket Queue View: Full search bar and multi-filter dropdown controls.",
                "Ticket Details View: Incident description, staff status controls, and activity timeline.",
                "Staff Login Portal: Clean authentication card with college demonstration credentials notice."
            ]
        },
        # Slide 9: Testing & Results
        {
            "title": "Testing, Results & Limitations",
            "subtitle": "Quality Assurance & Evaluation Metrics",
            "bullets": [
                "Testing Framework: Python standard unittest with isolated SQLite test databases.",
                "Test Coverage: 18 automated unit and integration tests covering routes, validation, and auth.",
                "Pass Rate: 100% (18 Passed, 0 Failed, 0 Skipped).",
                "Known Limitations:",
                "  - SQLite file concurrency limits simultaneous high-volume write operations.",
                "  - Status notifications are in-app only (no external SMS/SMTP gateway connected yet)."
            ]
        },
        # Slide 10: Conclusion & Future Scope
        {
            "title": "Conclusion & Future Scope",
            "subtitle": "Summary of Outcomes & Planned Version 2.0 Roadmap",
            "bullets": [
                "Conclusion: Successfully built a complete, secure, working IT Help Desk Ticket System.",
                "Achieved complete alignment with internship goals and academic evaluation criteria.",
                "Future Scope (Version 2.0):",
                "  - Automated SMTP email alerts to requesters upon ticket status changes.",
                "  - Diagnostic file and screenshot attachment uploads.",
                "  - SLA countdown timer for critical High-Priority tickets.",
                "Thank You! Open for Project Demonstration and Viva Panel Q&A."
            ]
        }
    ]

    for item in slides_data:
        slide = prs.slides.add_slide(blank_layout)

        # Title text box
        tb_title = slide.shapes.add_textbox(PptxInches(0.8), PptxInches(0.6), PptxInches(11.7), PptxInches(1.2))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = item["title"]
        p_title.font.size = PptxPt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = c_primary

        if "subtitle" in item:
            p_sub = tf_title.add_paragraph()
            p_sub.text = item["subtitle"]
            p_sub.font.size = PptxPt(16)
            p_sub.font.color.rgb = c_muted

        # Bullets text box
        tb_body = slide.shapes.add_textbox(PptxInches(0.8), PptxInches(2.0), PptxInches(11.7), PptxInches(4.8))
        tf_body = tb_body.text_frame
        tf_body.word_wrap = True

        for i, bullet in enumerate(item["bullets"]):
            p = tf_body.add_paragraph() if i > 0 else tf_body.paragraphs[0]
            p.text = bullet
            p.font.size = PptxPt(18)
            p.font.color.rgb = c_dark
            p.space_after = PptxPt(12)

    prs.save(output_path)
    print(f"Successfully generated PowerPoint presentation: {output_path}")


if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.abspath(__file__))
    word_file = os.path.join(base_dir, 'project_report.docx')
    pptx_file = os.path.join(base_dir, 'presentation.pptx')

    generate_word_report(word_file)
    generate_pptx_presentation(pptx_file)
